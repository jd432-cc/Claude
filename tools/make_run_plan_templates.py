#!/usr/bin/env python3
"""
Build the three Run Plan & Setup Log templates.

The existing taggers under tools/ re-tag a blank Word form that
somebody already drew. There is no blank form for these three, so this
one draws them and tags them in the same pass. That is the only
difference: the output is the same kind of artefact, tagged the same
way, validated by the same tools/validate_template.py.

Doing it in code rather than by hand is not a preference. Word splits a
run the moment you type near a tag, which turns {bestLap} into {best +
Lap} across two runs and breaks it silently: the export still succeeds
and the field comes out blank. Every tag below is written as exactly
one run, by construction.

Three rules the engine imposes, all of them learned the hard way and
recorded in _tools/report-builder/README.md:

  - a loop repeats everything between FOR and END-FOR, so each command
    needs a carrier row of its own; a row containing only a command is
    dropped by the engine;
  - a loop variable needs a $ prefix: {$p.label}, never {p.label};
  - a brace anywhere in the document is a command, so no stray ones.

Usage:  python3 tools/make_run_plan_templates.py
Then:   python3 tools/validate_template.py <out.docx> \
            _tools/run-plan/js/schemas/run-plan.js
"""
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = pathlib.Path(__file__).resolve().parent.parent / '_tools/run-plan/templates'

INK = RGBColor(0x0A, 0x0A, 0x0B)
MUTED = RGBColor(0x5A, 0x5A, 0x60)
SIGNAL = RGBColor(0x97, 0x00, 0x00)

BODY = 'Archivo'
FALLBACK = 'Arial'


# --------------------------------------------------------------------------
# Formatting helpers. Every one of them writes exactly one run.
# --------------------------------------------------------------------------

def run(paragraph, text, *, size=10, bold=False, italic=False,
        colour=INK, caps=False, spacing=None):
    r = paragraph.add_run(text)
    r.font.name = BODY
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = colour
    r.font.all_caps = caps
    if spacing is not None:
        # python-docx has no letter-spacing property; w:spacing is in
        # twentieths of a point and the element goes on by hand.
        el = OxmlElement('w:spacing')
        el.set(qn('w:val'), str(int(spacing * 20)))
        r._element.get_or_add_rPr().append(el)
    return r


def para(doc, text='', **kw):
    p = doc.add_paragraph()
    if text:
        run(p, text, **kw)
    return p


def eyebrow(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, text, size=7.5, bold=True, colour=MUTED, caps=True, spacing=1.8)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run(p, text, size=20, bold=True, colour=INK, caps=True)
    return p


def rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    pr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:color'), '0A0A0B')
    borders.append(bottom)
    pr.append(borders)
    return p


def guidance(doc, text):
    """A guidance line: kept in DRAFT, stripped from RELEASE.

    The bracket is three paragraphs so the command paragraphs can be
    removed with the line between them, which is how the engine's IF
    works on block content.
    """
    para(doc, '{IF showGuidance}', size=8, colour=MUTED)
    p = doc.add_paragraph()
    run(p, text, size=8.5, italic=True, colour=MUTED)
    para(doc, '{END-IF}', size=8, colour=MUTED)


def cell_text(cell, text, *, size=9, bold=False, colour=INK, caps=False, spacing=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if text:
        run(p, text, size=size, bold=bold, colour=colour, caps=caps, spacing=spacing)


def table(doc, cols, widths=None, header=None):
    t = doc.add_table(rows=0, cols=cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    if header:
        row = t.add_row()
        for cell, label in zip(row.cells, header):
            cell_text(cell, label, size=7.5, bold=True, colour=MUTED, caps=True, spacing=1.4)
    if widths:
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
        t._widths = widths
    return t


def row(t, values, **kw):
    r = t.add_row()
    for cell, value in zip(r.cells, values):
        cell_text(cell, value, **kw)
    widths = getattr(t, '_widths', None)
    if widths:
        for cell, w in zip(r.cells, widths):
            cell.width = Cm(w)
    return r


def carrier(t, command):
    """A row holding one command and nothing else.

    The engine drops it, which is why the command must not share a row
    with content: a FOR in a body row leaks a stray leading cell on
    every iteration.
    """
    r = t.add_row()
    cell_text(r.cells[0], command, size=8, colour=MUTED)
    for cell in r.cells[1:]:
        cell_text(cell, '')
    return r


def header_block(doc, title, pairs):
    eyebrow(doc, 'TheRacingData — Performance Engineering')
    heading(doc, title)
    rule(doc)
    t = table(doc, 4, widths=[3.2, 5.0, 3.2, 5.0])
    for i in range(0, len(pairs), 2):
        chunk = pairs[i:i + 2]
        while len(chunk) < 2:
            chunk.append(('', ''))
        row(t, [chunk[0][0], chunk[0][1], chunk[1][0], chunk[1][1]])
        for cell, bold in zip(t.rows[-1].cells, [True, False, True, False]):
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = bold
                    if bold:
                        r.font.color.rgb = MUTED
                        r.font.size = Pt(7.5)
                        r.font.all_caps = True
    para(doc)
    return t


def section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run(p, text, size=11, bold=True, colour=SIGNAL, caps=True, spacing=1.2)
    return p


def base_document():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = BODY
    style.font.size = Pt(10)
    for s in doc.sections:
        s.left_margin = Cm(1.6)
        s.right_margin = Cm(1.6)
        s.top_margin = Cm(1.4)
        s.bottom_margin = Cm(1.4)
    return doc


def footer(doc, note):
    para(doc)
    p = doc.add_paragraph()
    run(p, note, size=7.5, colour=MUTED)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


# --------------------------------------------------------------------------
# 1. Setup sheet — the current revision, print-ready for the car file
# --------------------------------------------------------------------------

def setup_sheet(path):
    doc = base_document()
    header_block(doc, 'Setup sheet', [
        ('Car', '{car.name}'), ('Class', '{car.class}'),
        ('Chassis', '{car.chassisNo}'), ('Reference', '{reportRef}'),
        ('Event', '{event.name}'), ('Venue', '{event.venue}'),
        ('Date', '{event.date}'), ('Session', '{event.sessionRef}'),
        ('Revision', '{rev}'), ('Based on', '{revBasedOn}'),
        ('Recorded', '{revTimestamp}'), ('Engineer', '{engineer}'),
        ('Note', '{revNote}'), ('Revisions on file', '{revisionCount}'),
    ])

    guidance(doc, 'One revision per sheet. A figure marked "computed" is derived from '
                  'the entries above it and is not enterable — if it disagrees with '
                  'this sheet, this sheet is the one that is wrong.')

    section(doc, 'Setup')
    t = table(doc, 4, widths=[7.6, 3.4, 2.4, 3.0],
              header=['Parameter', 'Value', 'Unit', 'Source'])
    carrier(t, '{FOR g IN setupGroups}')
    row(t, ['{$g.group}', '', '', ''], size=9, bold=True, caps=True, spacing=1.2)
    carrier(t, '{FOR p IN $g.params}')
    row(t, ['{$p.label}', '{$p.value}', '{$p.unit}', '{$p.derived}'])
    carrier(t, '{END-FOR p}')
    carrier(t, '{END-FOR g}')

    footer(doc, 'TheRacingData — setup-sheet-v1.0. Every derived figure on this sheet '
                'is computed from the entries beside it. History is never rewritten: a '
                'change taken back is a new revision, not a deleted one.')
    doc.save(path)
    return path


# --------------------------------------------------------------------------
# 2. Run plan — the session's runs, one row per run, for the pit wall
# --------------------------------------------------------------------------

def run_plan(path):
    doc = base_document()
    for s in doc.sections:
        # Landscape: a run row is wide and a folded sheet is not read.
        s.orientation = 1
        s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = Cm(1.2)
        s.right_margin = Cm(1.2)

    header_block(doc, 'Run plan', [
        ('Car', '{car.name}'), ('Reference', '{reportRef}'),
        ('Event', '{event.name}'), ('Session', '{event.sessionRef}'),
        ('Date', '{event.date}'), ('Engineer', '{engineer}'),
        ('Runs', '{runCount}'), ('Best lap', '{bestLap}'),
        ('Kept', '{keptCount}'), ('Reverted', '{revertedCount}'),
        ('Inconclusive', '{inconclusiveCount}'), ('Current revision', '{currentRev}'),
    ])

    guidance(doc, 'One variable per run. A run that changed more than one thing cannot '
                  'attribute its result to any of them, and is marked in the last '
                  'column when it was done deliberately.')

    section(doc, 'Runs')
    t = table(doc, 8, widths=[1.1, 4.4, 2.2, 1.6, 4.6, 4.4, 4.4, 2.4],
              header=['#', 'Objective', 'Driver', 'Rev', 'Change', 'Expected',
                      'Measured', 'Verdict'])
    carrier(t, '{FOR r IN runRows}')
    row(t, ['{$r.n}', '{$r.objective}', '{$r.driver}', '{$r.setupRev}',
            '{$r.change}', '{$r.expected}', '{$r.measured}', '{$r.verdict}'])
    row(t, ['', '{$r.rationale}', '{$r.tyreSet}', '{$r.laps}',
            '{$r.multiChange}', '{$r.bestLap}', '{$r.notes}', '{$r.fuel}'],
        size=8, colour=MUTED)
    carrier(t, '{END-FOR r}')

    footer(doc, 'TheRacingData — run-plan-v1.0. Second line per run: rationale, tyre '
                'set, laps, multi-change note, best lap, notes, fuel. A verdict of keep '
                'or revert requires a measured result.')
    doc.save(path)
    return path


# --------------------------------------------------------------------------
# 3. Setup diff — A against B, with the reconciliation table
# --------------------------------------------------------------------------

def setup_diff(path):
    doc = base_document()
    header_block(doc, 'Setup diff', [
        ('Car', '{car.name}'), ('Reference', '{reportRef}'),
        ('Event', '{event.name}'), ('Session', '{event.sessionRef}'),
        ('Revision A', '{compare.a}'), ('Revision B', '{compare.b}'),
        ('A note', '{compare.aNote}'), ('B note', '{compare.bNote}'),
    ])

    guidance(doc, 'Unchanged parameters are omitted. A numeric parameter carries a '
                  'signed delta; an enumerated one carries from and to, because the '
                  'difference between two compounds is not a number.')

    section(doc, 'Changed')
    t = table(doc, 6, widths=[3.0, 5.0, 2.4, 2.4, 2.4, 2.4],
              header=['Group', 'Parameter', 'From', 'To', 'Delta', 'Delta %'])
    carrier(t, '{FOR d IN diffRows}')
    row(t, ['{$d.group}', '{$d.label}', '{$d.from}', '{$d.to}',
            '{$d.delta}', '{$d.deltaPct}'])
    carrier(t, '{END-FOR d}')

    section(doc, 'Consequences')
    guidance(doc, 'Derived figures that moved because of the changes above. They were '
                  'not set; they were caused.')
    c = table(doc, 4, widths=[3.0, 6.6, 3.8, 3.8],
              header=['Group', 'Figure', 'From', 'To'])
    carrier(c, '{FOR v IN derivedRows}')
    row(c, ['{$v.group}', '{$v.label}', '{$v.from}', '{$v.to}'])
    carrier(c, '{END-FOR v}')

    section(doc, 'Reconciliation')
    guidance(doc, 'Undeclared is what moved and was not declared. Not applied is what '
                  'was declared and did not move. Both are common and both are silent '
                  'failures on paper.')
    r = table(doc, 5, widths=[1.4, 4.4, 4.4, 4.0, 3.0],
              header=['Run', 'Declared', 'Undeclared', 'Not applied', 'Verdict'])
    carrier(r, '{FOR c IN reconRows}')
    row(r, ['{$c.run}', '{$c.declared}', '{$c.undeclared}',
            '{$c.notApplied}', '{$c.verdict}'])
    carrier(r, '{END-FOR c}')

    footer(doc, 'TheRacingData — setup-diff-v1.0.')
    doc.save(path)
    return path


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    for fn, name in [(setup_sheet, 'setup-sheet-v1.0.docx'),
                     (run_plan, 'run-plan-v1.0.docx'),
                     (setup_diff, 'setup-diff-v1.0.docx')]:
        path = fn(OUT / name)
        print(f'  wrote  {path.relative_to(OUT.parent.parent.parent)}')
